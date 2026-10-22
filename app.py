import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from pathlib import Path

from flask import Flask, render_template, request, send_file
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "template" / "sih2026_nomination_template.docx"

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024

COLUMN_WIDTHS_IN = [
    1.00,  # Team Leader / Team Member
    1.20,  # Name
    0.65,  # Gender
    1.90,  # Email id
    1.05,  # Mobile no.
    0.65,  # Stream
    0.85,  # Academic Year
]

EMAIL_RE = re.compile(r"^[A-Za-z]+\.[A-Za-z]+[0-9]{2}@bmsce\.ac\.in$", re.I)
STREAM_RE = re.compile(r"^[A-Z]{2,6}$")
YEAR_VALUES = {"1st", "2nd", "3rd", "4th"}


def set_run_font(run, size_pt=11, bold=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), "Times New Roman")
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_cell_margins(cell, top=0, start=144, bottom=0, end=144):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_no_wrap(cell, enabled=True):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:noWrap"))
    if enabled and existing is None:
        tc_pr.append(OxmlElement("w:noWrap"))
    elif not enabled and existing is not None:
        tc_pr.remove(existing)


def set_cell_text(cell, text, alignment, size_pt=11, bold=False, no_wrap=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, 0, 144, 0, 144)
    set_no_wrap(cell, no_wrap)
    for paragraph in cell.paragraphs:
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size_pt, bold)


def make_output_stem(team_name, team_id):
    team_name = " ".join(str(team_name).split())
    safe_name = re.sub(r"\s+", "_", team_name)
    safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", safe_name)
    safe_name = re.sub(r"_+", "_", safe_name).strip(" ._") or "Team"
    return f"BMS_SIH2026_{team_id}_{safe_name}"


def validate_submission(team_name, team_id, people):
    errors = []
    team_name = " ".join(str(team_name).split())
    team_id = str(team_id).strip()

    if not team_name:
        errors.append("Team Name: enter your team name.")
    elif len(team_name) > 80:
        errors.append("Team Name: maximum 80 characters.")
    elif any(ch in '<>:/\\|?*"' or ord(ch) < 32 for ch in team_name):
        errors.append("Team Name: avoid / \\ : * ? \" < > |.")

    if not re.fullmatch(r"[0-9]+", team_id):
        errors.append("Team ID: enter numbers only. Example: 65.")

    seen_emails = {}
    seen_mobiles = {}

    for index, person in enumerate(people):
        role = "Team Leader" if index == 0 else f"Team Member {index}"
        name = " ".join(str(person.get("name", "")).split())
        gender = str(person.get("gender", "")).strip().upper()
        email = str(person.get("email", "")).strip()
        mobile = str(person.get("mobile", "")).strip()
        stream = str(person.get("stream", "")).strip().upper()
        year = str(person.get("year", "")).strip()

        if not name:
            errors.append(f"{role}: enter the full name as per government records.")
        if gender not in {"M", "F"}:
            errors.append(f"{role}: Gender must be M or F.")
        if not EMAIL_RE.fullmatch(email):
            errors.append(f"{role}: enter a valid BMSCE email such as name.aa00@bmsce.ac.in.")
        elif email.lower() in seen_emails:
            errors.append(f"{role}: this email is already used by {seen_emails[email.lower()] }.")
        else:
            seen_emails[email.lower()] = role

        if not re.fullmatch(r"[0-9]{10}", mobile):
            errors.append(f"{role}: mobile number must contain exactly 10 digits with no spaces or symbols.")
        elif mobile in seen_mobiles:
            errors.append(f"{role}: this mobile number is already used by {seen_mobiles[mobile]}.")
        else:
            seen_mobiles[mobile] = role

        if not STREAM_RE.fullmatch(stream):
            errors.append(f"{role}: stream must be an uppercase code such as CSE, ECE, ISE, EEE or AIML.")
        if year not in YEAR_VALUES:
            errors.append(f"{role}: Academic Year must be 1st, 2nd, 3rd or 4th.")

    return errors


def generate_files(team_name, team_id, people, workdir):
    errors = validate_submission(team_name, team_id, people)
    if errors:
        raise ValueError("\n".join(f"• {e}" for e in errors))

    team_name = " ".join(str(team_name).split())
    team_id = str(team_id).strip()
    doc = Document(str(TEMPLATE_PATH))

    for section in doc.sections:
        section.top_margin = Cm(4)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, run.font.size.pt if run.font.size else 12, run.bold)

    team_paragraph = next((p for p in doc.paragraphs if p.text.strip().lower().startswith("team")), None)
    if team_paragraph is None:
        raise RuntimeError("Could not find the Team Name line in the template.")
    team_paragraph.clear()
    run = team_paragraph.add_run(f"Team: {team_name}")
    set_run_font(run, 12, True)

    if not doc.tables:
        raise RuntimeError("No participant table was found in the template.")
    table = doc.tables[0]
    if len(table.rows) != 7 or len(table.columns) != 7:
        raise RuntimeError("The participant table must contain 7 rows and 7 columns.")
    table.autofit = False

    for ci, width in enumerate(COLUMN_WIDTHS_IN):
        table.columns[ci].width = Inches(width)
        for row in table.rows:
            row.cells[ci].width = Inches(width)

    headers = ["", "Name", "Gender", "Email id", "Mobile no.", "Stream", "Academic\nYear"]
    for ci, header in enumerate(headers):
        set_cell_text(table.cell(0, ci), header, WD_ALIGN_PARAGRAPH.CENTER, 8, True, False)

    roles = ["Team\nLeader"] + ["Team\nMember"] * 5
    for ri in range(1, 7):
        person = people[ri - 1]
        set_cell_text(table.cell(ri, 0), roles[ri - 1], WD_ALIGN_PARAGRAPH.CENTER, 11, False, False)
        set_cell_text(table.cell(ri, 1), " ".join(person["name"].split()), WD_ALIGN_PARAGRAPH.LEFT, 11, False, False)
        set_cell_text(table.cell(ri, 2), person["gender"], WD_ALIGN_PARAGRAPH.CENTER, 11, False, True)
        set_cell_text(table.cell(ri, 3), person["email"].strip(), WD_ALIGN_PARAGRAPH.LEFT, 11, False, True)
        set_cell_text(table.cell(ri, 4), person["mobile"].strip(), WD_ALIGN_PARAGRAPH.CENTER, 11, False, True)
        set_cell_text(table.cell(ri, 5), person["stream"].strip().upper(), WD_ALIGN_PARAGRAPH.CENTER, 11, False, True)
        set_cell_text(table.cell(ri, 6), person["year"].strip(), WD_ALIGN_PARAGRAPH.CENTER, 11, False, True)

    for ri, row in enumerate(table.rows):
        row.height = Inches(0.36 if ri == 0 else 0.40)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, run.font.size.pt if run.font.size else 11, run.bold)

    stem = make_output_stem(team_name, team_id)
    docx_path = workdir / f"{stem}.docx"
    pdf_path = workdir / f"{stem}.pdf"
    txt_path = workdir / f"{stem}.txt"
    zip_path = workdir / f"{stem}.zip"

    doc.save(str(docx_path))

    profile = workdir / f"lo-profile-{uuid.uuid4().hex}"
    profile.mkdir()
    result = subprocess.run([
        "libreoffice", f"-env:UserInstallation={profile.as_uri()}",
        "--headless", "--convert-to", "pdf", "--outdir", str(workdir), str(docx_path)
    ], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "PDF conversion failed.")

    lines = [
        "SIH 2026 Team Nomination - Input Information",
        "=" * 48,
        f"Team Name: {team_name}",
        f"Team ID: BMS_SIH2026_{team_id}",
        "",
    ]
    for i, p in enumerate(people):
        role = "Team Leader" if i == 0 else f"Team Member {i}"
        lines += [
            role,
            "-" * len(role),
            f"Name: {' '.join(p['name'].split())}",
            f"Gender: {p['gender']}",
            f"College Email: {p['email'].strip()}",
            f"Mobile: {p['mobile'].strip()}",
            f"Stream: {p['stream'].strip().upper()}",
            f"Academic Year: {p['year'].strip()}",
            "",
        ]
    txt_path.write_text("\n".join(lines), encoding="utf-8")

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(docx_path, docx_path.name)
        z.write(pdf_path, pdf_path.name)
        z.write(txt_path, txt_path.name)

    return docx_path, pdf_path, txt_path, zip_path


@app.get("/")
def index():
    return render_template("index.html")


@app.post("/generate")
def generate():
    team_name = request.form.get("team_name", "")
    team_id = request.form.get("team_id", "")
    people = []
    for i in range(6):
        people.append({
            "name": request.form.get(f"name_{i}", ""),
            "gender": request.form.get(f"gender_{i}", ""),
            "email": request.form.get(f"email_{i}", ""),
            "mobile": request.form.get(f"mobile_{i}", ""),
            "stream": request.form.get(f"stream_{i}", ""),
            "year": request.form.get(f"year_{i}", ""),
        })

    errors = validate_submission(team_name, team_id, people)
    if errors:
        return render_template("index.html", errors=errors, form=request.form), 400

    with tempfile.TemporaryDirectory(prefix="sih-") as tmp:
        workdir = Path(tmp)
        try:
            _, _, _, zip_path = generate_files(team_name, team_id, people, workdir)
            # copy outside TemporaryDirectory so send_file can read it after function returns
            persistent = Path(tempfile.mkdtemp(prefix="sih-download-")) / zip_path.name
            shutil.copy2(zip_path, persistent)
        except Exception as exc:
            return render_template("index.html", errors=[f"Error while generating files: {exc}"], form=request.form), 500

    return send_file(persistent, as_attachment=True, download_name=persistent.name, mimetype="application/zip")


@app.get("/health")
def health():
    return {"status": "ok"}


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "10000"))
    app.run(host="0.0.0.0", port=port, debug=False)
